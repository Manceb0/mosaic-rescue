#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Empaqueta los informes de entrega (Markdown + PDF) con el formulario de metodos dentro.

    python src/submission/export_reports.py            # dependencias: openpyxl, markdown, playwright (Edge o Chromium)

POR QUE
    La plataforma de Track 1 solo tiene dos casillas: el CSV y UN informe (PDF o Markdown). La
    plantilla oficial de metodos (Excel) se pide "para organizar" ese informe, asi que sus
    respuestas tienen que viajar DENTRO del informe o el jurado no las ve. Este script:

    1. el informe maestro manda: las respuestas viven en el apendice de reports/TRACK{1,2}_REPORT.md,
       entre marcadores, y se ESCRIBEN en el Excel de entrega
       (work/submission/subir/<equipo>_methods_description.xlsx), nunca al reves. Hasta 2026-10-07 el
       sentido era Excel -> informe; con el apendice corregido despues (HLA-DRB1, Exomiser «externo»,
       EXP-029), ejecutar el exportador antiguo habria revertido esas correcciones con el Excel del
       26-sep. Una pregunta del Excel que no aparezca en el apendice detiene el proceso;
    2. escribe la copia Markdown de entrega con enlaces absolutos al repositorio;
    3. escribe un PDF con las figuras INCRUSTADAS, porque mientras el repositorio sea privado
       las imagenes enlazadas no se ven.

    No toca el CSV ni lee datos del paciente.
"""
from __future__ import annotations
import os, re, shutil, sys, pathlib

TEAM = "mosaic-rescue"
REPO_URL = "https://github.com/Manceb0/mosaic-rescue/blob/main/"
OUT = "work/submission/subir"
XLSX = os.path.join(OUT, f"{TEAM}_methods_description.xlsx")
EXTRA_OUT = os.environ.get("MVA_DELIVERY_DIR")  # copia opcional fuera del repo (p. ej. Descargas)
START, END = "<!-- methods-form:start -->", "<!-- methods-form:end -->"
REPORTS = {1: ("reports/TRACK1_REPORT.md", "Track 1 methods", range(7, 20)),
           2: ("reports/TRACK2_REPORT.md", "Track 2 methods", range(7, 18))}

CSS = """body{font-family:Cambria,Georgia,'Times New Roman',serif;font-size:10.5pt;line-height:1.58;color:#161616;
  max-width:164mm;margin:auto}
h1{font-size:21pt;font-weight:700;margin:0 0 6pt;line-height:1.2}
h2{font-size:13pt;font-weight:700;margin:26pt 0 8pt}
h3{font-size:11pt;font-weight:700;font-style:italic;margin:16pt 0 5pt}
p{margin:0 0 10pt;text-align:left}li{margin:0 0 5pt}ul,ol{margin:0 0 10pt}
table{border-collapse:collapse;width:100%;font-size:9pt;line-height:1.35;margin:8pt 0 10pt;
  border-top:1.2pt solid #111;border-bottom:1.2pt solid #111}
th{font-weight:700;text-align:left;border-bottom:.6pt solid #111;padding:3pt 6pt;vertical-align:bottom}
td{padding:3pt 6pt;vertical-align:top;border:none}
code{font-family:inherit;font-size:inherit;background:none;padding:0}
pre{font-family:Consolas,'Courier New',monospace;font-size:8.5pt;margin:4pt 0 8pt 12pt;white-space:pre-wrap}
pre code{font-family:inherit}
blockquote{margin:6pt 18pt;padding:0;color:#333;font-style:italic}
img{max-width:100%;display:block;margin:0 auto}
figure{margin:16pt 0 18pt;clear:both;break-inside:avoid}
figure img{max-height:95mm;max-width:100%;width:auto;margin:0 auto 7pt}figure img+img{margin-top:6pt}
.pair{display:grid;grid-template-columns:1fr 54%;column-gap:16pt;align-items:start;margin:6pt 0 16pt;
  break-inside:avoid}
.pair.left{grid-template-columns:54% 1fr}.pair>p{margin:0}.pair figure{margin:2pt 0 0}
figure.side img{width:100%;max-height:none}figure.mid{width:72%;margin:16pt auto 18pt}
p.cap{font-size:8.6pt;line-height:1.42;color:#4a4a4a;margin:0}p.cap b{color:#161616}
.toc-h{font-size:9.5pt;font-weight:700;margin:14pt 0 4pt}
ol.toc{columns:2;column-gap:20pt;list-style:none;padding:0;margin:0 0 6pt;font-size:9pt;line-height:1.55}
ol.toc li{margin:0;break-inside:avoid}ol.toc a{color:#161616}
a{color:#1d3f6e;text-decoration:none}hr{display:none}
h1,h2,h3{break-after:avoid}
img,tr,blockquote,pre{break-inside:avoid}
p:has(+ table){break-after:avoid}
thead:not(:has(th:not(:empty))){display:none}"""
# Formato de artículo, no de interfaz (2026-10-07, a petición del usuario): texto con serifa, títulos sin subrayar,
# tablas de tres rayas sin rejilla ni fondos, nombres y rutas sin «fichas» grises, sin reglas horizontales. Las
# últimas cuatro reglas vienen de revisar los PDF a ojo: un título solo al pie de página con su figura en la
# siguiente, una frase separada de su tabla y una fila de cabecera vacía en las tablas de dos columnas.


def appendix_answers(md: str) -> dict:
    """Pregunta -> respuesta, leidas del apendice del informe maestro (entre START y END)."""
    block = md[md.index(START):md.index(END)]
    parts = re.split(r"^\*\*(.+?)\*\*\s*$", block, flags=re.M)
    return {q.strip(): a.strip() for q, a in zip(parts[1::2], parts[2::2])}


def write_form(wb, sheet: str, rows, answers: dict) -> int:
    """Escribe en el Excel las respuestas del informe; falla si una pregunta no tiene respuesta."""
    ws, n = wb[sheet], 0
    for r in rows:
        q = ws.cell(r, 1).value
        if not q:
            continue
        key = str(q).strip()
        if key not in answers:
            raise SystemExit(f"{sheet}: la pregunta de la fila {r} no esta en el apendice del informe: {key[:60]}")
        ws.cell(r, 2).value = answers[key]
        n += 1
    return n


def absolute_links(md: str) -> str:
    def fix(m):
        bang, txt, url = m.groups()
        if url.startswith(("http", "#", "mailto", "file:")):
            return m.group(0)
        p = re.sub(r"^(\.\./)+", "", url)
        return f"{bang}[{txt}]({REPO_URL}{p}{'?raw=true' if bang else ''})"
    return re.sub(r"(!?)\[([^\]]*)\]\(([^)\s]+)\)", fix, md)


def local_images(md: str, base: str) -> str:
    def fix(m):
        txt, url = m.groups()
        if url.startswith("http"):
            return m.group(0)
        return f"![{txt}]({pathlib.Path(base, url).resolve().as_uri()})"
    md = re.sub(r"!\[([^\]]*)\]\(([^)\s]+)\)", fix, md)
    return absolute_links(md)  # los enlaces normales, al repositorio


def figure_block(m, side: list) -> str:
    imgs, n, cap = re.sub(r"</?p>", "", m.group(1)).strip(), m.group(2), m.group(3)
    cls = ""
    if "drawn with AI image generation" in cap and imgs.count("<img") == 1:
        cls = ' class="side"'
    return f'<figure{cls}>{imgs}<p class="cap"><b>Figure {n}.</b> {cap}</p></figure>'


def pair_figures(body: str) -> str:
    """Cada esquema se empareja con el párrafo anterior (o, si no vale, el siguiente) en dos columnas, alternando
    el lado. Un párrafo de 250 a 900 caracteres llena la columna sin dejar hueco. Sin párrafo adecuado, el
    esquema va centrado a menor ancho, nunca flotando solo junto a un título o una tabla."""
    para = r"<p>(?:(?!</p>).)*</p>"
    fits = lambda h: 250 <= len(re.sub(r"<[^>]+>", "", h)) <= 900
    figs = list(re.finditer(r'<figure class="side">.*?</figure>', body, flags=re.S))
    n = len(figs)
    for k, f in enumerate(reversed(figs)):
        side = "" if (n - 1 - k) % 2 == 0 else " left"
        a, b, fig = f.start(), f.end(), f.group(0)
        prev = re.search(rf"({para})\s*$", body[:a], flags=re.S)
        nxt = re.match(rf"\s*({para})", body[b:], flags=re.S)
        if prev and fits(prev.group(1)):
            a, text = prev.start(1), prev.group(1)
        elif nxt and fits(nxt.group(1)):
            b, text = b + nxt.end(1), nxt.group(1)
        else:
            body = body[:a] + fig.replace('class="side"', 'class="mid"') + body[b:]
            continue
        cells = text + fig if not side else fig + text
        body = body[:a] + f'<div class="pair{side}">{cells}</div>' + body[b:]
    return body


def to_pdf(md: str, pdf: str, title: str):
    import markdown
    # «6 321» no debe partirse al final de una línea: espacio que no se rompe entre grupos de miles.
    md = re.sub(r"(?<=\d) (?=\d{3}(?!\d))", " ", md)
    from playwright.sync_api import sync_playwright
    body = markdown.markdown(md, extensions=["tables", "fenced_code", "toc"])
    # Índice enlazado tras el resumen: la página 1 abre con el hallazgo y el jurado salta luego a lo que busca.
    heads = re.findall(r'<h2 id="([^"]+)">(.*?)</h2>', body)
    toc = "".join(f'<li><a href="#{i}">{t}</a></li>' for i, t in heads)
    first = body.index("<h2", body.index("<h2") + 1)
    body = body[:first] + f'<p class="toc-h">Contents</p><ol class="toc">{toc}</ol>' + body[first:]
    # Los esquemas de las ilustraciones van a un lado con el texto alrededor, alternando derecha e izquierda,
    # como en una revista. Las figuras de datos, anchas y con letra pequeña, quedan a todo el ancho y sin marco.
    side = [0]
    body = re.sub(r"((?:<p><img [^>]*></p>\s*)+)<p><em>Figure (\d+)\. (.*?)</em></p>",
                  lambda m: figure_block(m, side), body, flags=re.S)
    body = pair_figures(body)
    html = f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title><style>{CSS}</style></head><body>" \
           + body + "</body></html>"
    tmp = pathlib.Path(pdf).with_suffix(".tmp.html"); tmp.write_text(html, encoding="utf-8")
    with sync_playwright() as p:
        try:
            b = p.chromium.launch(channel="msedge", headless=True)
        except Exception:
            b = p.chromium.launch(headless=True)
        pg = b.new_page(); pg.goto(tmp.resolve().as_uri()); pg.wait_for_load_state("networkidle")
        pg.pdf(path=pdf, format="A4", margin={"top": "16mm", "bottom": "16mm", "left": "14mm", "right": "14mm"},
               print_background=True, outline=True, tagged=True, display_header_footer=True, header_template="<span></span>",
               footer_template="<div style='font-size:8px;width:100%;text-align:center;color:#777'>"
                               f"{title}, page <span class='pageNumber'></span> of <span class='totalPages'></span></div>")
        b.close()
    tmp.unlink()


def main() -> int:
    import openpyxl
    assert os.path.exists(XLSX), f"falta la plantilla del formulario: {XLSX}"
    wb = openpyxl.load_workbook(XLSX)
    try:
        import markdown  # noqa: F401  (solo para el PDF)
        pdf_ok = True
    except ImportError:
        pdf_ok = False
        print("AVISO: falta el paquete `markdown`; se escriben .md y .xlsx, NO el .pdf")
    for n, (src, sheet, rows) in REPORTS.items():
        md = open(src, encoding="utf-8").read()
        k = write_form(wb, sheet, rows, appendix_answers(md))
        stem = os.path.join(OUT, f"{TEAM}_track{n}_report")
        open(stem + ".md", "w", encoding="utf-8", newline="\n").write(absolute_links(md))
        exts = [".md"]
        if pdf_ok:
            to_pdf(local_images(md, os.path.dirname(src)), stem + ".pdf", f"MOSAIC RESCUE, Track {n} report")
            exts.append(".pdf")
        for ext in exts:
            if EXTRA_OUT:
                shutil.copy(stem + ext, EXTRA_OUT)
        print(f"track {n}: {k} respuestas del apendice de {src} escritas en el Excel; "
              f"{', '.join(stem + e for e in exts)}")
    wb.save(XLSX)
    if EXTRA_OUT:
        shutil.copy(XLSX, EXTRA_OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
